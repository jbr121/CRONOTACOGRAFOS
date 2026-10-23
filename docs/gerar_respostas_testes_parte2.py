"""Respostas da atividade de Testes - parte 2 (linguagem simples)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1B, 0x4F, 0x72)
PRETO = RGBColor(0x1C, 0x28, 0x33)


def _font(run, size: int = 12, bold: bool = False, italic: bool = False, color=PRETO) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _p(doc: Document, texto: str, *, size: int = 12, bold: bool = False,
       italic: bool = False, center: bool = False, after: int = 8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    _font(run, size=size, bold=bold, italic=italic)


def _titulo(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texto)
    _font(run, size=13, bold=True, color=AZUL)


def montar(saida: Path) -> None:
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    _p(doc, "Atividade — Testes de software (parte 2)", size=18, bold=True, center=True, after=4)
    _p(doc, "Respostas em linguagem simples", size=12, italic=True, center=True, after=14)

    _titulo(doc, "1. Em uma verificação de produtos ou documentos, o que se busca?")
    _p(
        doc,
        "Em uma verificação nós procuramos erros/defeitos nos produtos ou documentos "
        "(requisitos, diagramas, código, planos). A ideia é achar problemas cedo, "
        "antes de o sistema rodar para o usuário, para ficar mais barato de corrigir.",
    )

    _titulo(doc, "2. Como diminuir defeitos nas fases iniciais do desenvolvimento?")
    _p(
        doc,
        "Algumas formas: conversar bem com o cliente e registrar requisitos claros; "
        "revisar requisitos e modelos em grupo; fazer protótipos de tela cedo para "
        "validar entendimento; usar padrões e checklists; e fazer revisão de código "
        "desde o começo, não só no fim do projeto.",
    )

    _titulo(doc, "3. \"Estamos construindo o produto corretamente\" x \"o produto certo\"")
    _p(
        doc,
        "\"Estamos construindo o produto corretamente?\" quer dizer: estamos seguindo "
        "o que foi especificado? O programa está de acordo com os requisitos, "
        "diagramas e regras definidos. Já \"Estamos construindo o produto certo?\" "
        "pergunta se o que estamos fazendo realmente resolve o problema do cliente, "
        "se ele vai ficar satisfeito usando o sistema.",
    )

    _titulo(doc, "4. Três produtos que devem passar por técnica estática de V&V")
    _p(
        doc,
        "Exemplos de produtos que vale muito a pena revisar (verificação estática): "
        "documento de requisitos, diagrama de casos de uso/fluxo e código fonte de "
        "um módulo crítico. Critério: são itens que guiam o resto do trabalho; se "
        "estiverem errados, o erro se espalha. E é bem mais barato corrigir nessa fase.",
    )

    _titulo(doc, "5. Como organizar os testes dinâmicos para serem mais efetivos?")
    _p(
        doc,
        "Mesmo precisando do programa pronto, dá para planejar antes: definir plano "
        "de testes, casos e dados; começar por testes de unidade, depois integração e "
        "sistema; priorizar o que tem mais risco para o negócio; automatizar o que se "
        "repete; e registrar defeitos para re-testar depois das correções.",
    )

    _titulo(doc, "6. Diferença entre técnicas estáticas e dinâmicas de V&V")
    _p(
        doc,
        "Técnicas estáticas não executam o programa: são leituras, revisões, análise "
        "de código, inspeção de requisitos. Elas verificam clareza, consistência, "
        "erros de lógica óbvios e padrões. Técnicas dinâmicas executam o programa, "
        "ou seja, são os testes. Elas verificam se, com certas entradas, o sistema "
        "se comporta como esperado e se não quebra em tempo de execução.",
    )

    _titulo(doc, "7. Exemplo de área em que os testes devem ser mais rigorosos")
    _p(
        doc,
        "Um exemplo é a área financeira/bancária. Qualquer erro pode gerar perda de "
        "dinheiro, fraude ou problema legal. Por isso os testes de operações como "
        "transferência, pagamento e cálculo de juros precisam ser muito mais "
        "rigorosos, com mais cenários, revisões e automação.",
    )

    _titulo(doc, "8. O que são testes caixa-branca e caixa-preta?")
    _p(
        doc,
        "Caixa-preta: o testador olha só entradas e saídas, sem ver o código. "
        "Exemplo: testar uma tela de login com usuários válidos e inválidos e "
        "ver se o sistema deixa entrar ou mostra erro. Caixa-branca: o testador "
        "usa o conhecimento do código (ifs, laços, caminhos) para montar os testes. "
        "Exemplo: criar casos para passar por todos os ramos de uma função que "
        "calcula desconto.",
    )

    _titulo(doc, "9. Testes de validação só verificam atendimento aos requisitos?")
    _p(
        doc,
        "A ideia principal da validação é sim conferir se o sistema atende aos "
        "requisitos do cliente. Mas vai além: também verifica se o cliente realmente "
        "aceita aquele comportamento no uso real. Às vezes o requisito estava mal "
        "escrito ou incompleto; na validação, o usuário testa o sistema e diz se o "
        "produto faz sentido para ele.",
    )

    _titulo(doc, "10. Classes de equivalência e valores-limite para a nota de 0 a 10")
    _p(
        doc,
        "Para o campo nota (0 a 10), as classes de equivalência são: valores válidos "
        "entre 0 e 10; valores menores que 0; valores maiores que 10; e valores "
        "não numéricos. Os valores-limite importantes são: -1, 0, 1, 9, 10 e 11.",
    )

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent / "Respostas_Atividade_Testes_parte2.docx"
    montar(destino)
    print(destino)

