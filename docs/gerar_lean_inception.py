"""Gera o cartaz Lean Inception do CronotaScan (PGI) — curto e fácil de postar."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1B, 0x4F, 0x72)
PRETO = RGBColor(0x1C, 0x28, 0x33)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _font(run, size=12, bold=False, italic=False, color=PRETO) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _sombra(cell, hex_cor: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_cor)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _bordas(table) -> None:
    tblPr = table._tbl.tblPr if table._tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "1B4F72")
        borders.append(el)
    tblPr.append(borders)


def _p(doc, texto, *, size=12, bold=False, italic=False, center=False, after=8, before=0) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 1.15
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    _font(run, size=size, bold=bold, italic=italic)


def _titulo(doc, texto) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(texto)
    _font(run, size=16, bold=True, color=AZUL)


def _subtitulo(doc, texto) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(texto)
    _font(run, size=13, bold=True, color=AZUL)


def _celula(cell, texto, *, header=False, fill=None, center=False, size=11) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto)
    if header:
        _font(run, size=12, bold=True, color=BRANCO)
        _sombra(cell, "1B4F72")
    else:
        _font(run, size=size, bold=False, color=PRETO)
        if fill:
            _sombra(cell, fill)


def _tabela_2x2(doc, titulos: list[str], colunas: list[list[str]], fills: list[str]) -> None:
    table = doc.add_table(rows=1 + max(len(c) for c in colunas), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordas(table)
    for i, titulo in enumerate(titulos):
        _celula(table.rows[0].cells[i], titulo, header=True, center=True)
    n = max(len(c) for c in colunas)
    for r in range(n):
        for c in range(2):
            texto = colunas[c][r] if r < len(colunas[c]) else ""
            _celula(table.rows[r + 1].cells[c], texto, fill=fills[c], size=11)
    doc.add_paragraph()


def montar(saida: Path) -> None:
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    _p(doc, "PROJETO INTEGRADOR — UNIVALI", size=11, bold=True, center=True, after=2)
    _p(doc, "Sistemas para Internet", size=11, center=True, after=10)
    _p(doc, "CronotaScan", size=26, bold=True, center=True, after=2)
    _p(doc, "Lean Inception — visão do produto", size=14, italic=True, center=True, after=4)
    _p(doc, "José Eduardo  ·  Márcio Fossa", size=12, center=True, after=18)

    # 1. Visão
    _titulo(doc, "1. Visão do produto")
    _p(
        doc,
        "A visão abaixo segue o formato da Lean Inception (CAROLI): uma frase "
        "simples que explica para quem é o produto, qual problema resolve e o "
        "que o diferencia.",
        after=10,
    )
    _subtitulo(doc, "Pitch (visão em uma frase)")

    blocos = [
        ("Para", "gestores de frota, auditores e operadores que conferem discos de cronotacógrafo"),
        ("Cujo problema é", "ler o papel circular de 24 horas no olho, de forma lenta, subjetiva e difícil de guardar como prova"),
        ("O", "CronotaScan"),
        ("É um", "leitor visual de discos analógicos, usado no computador local"),
        ("Que", "transforma a foto do disco em curva de velocidade, alertas e laudo"),
        ("Diferente de", "a leitura só no papel, da planilha manuscrita ou de equipamentos caros de oficina"),
        ("Nosso produto", "padroniza a conferência: placa + foto → dashboard + histórico + PDF, sem precisar de nuvem ou de login"),
    ]
    table = doc.add_table(rows=len(blocos), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordas(table)
    for i, (rotulo, texto) in enumerate(blocos):
        _celula(table.rows[i].cells[0], rotulo, fill="1B4F72" if True else None, size=11)
        table.rows[i].cells[0].paragraphs[0].clear()
        p = table.rows[i].cells[0].paragraphs[0]
        run = p.add_run(rotulo)
        _font(run, size=11, bold=True, color=BRANCO)
        _sombra(table.rows[i].cells[0], "1B4F72")
        _celula(table.rows[i].cells[1], texto, fill="EAF2F8", size=12)
        table.rows[i].cells[0].width = Cm(4.2)
        table.rows[i].cells[1].width = Cm(12.6)
    doc.add_paragraph()

    _subtitulo(doc, "Visão em texto corrido")
    _p(
        doc,
        "Para quem precisa conferir discos de cronotacógrafo, o CronotaScan é um "
        "sistema web local que lê a foto do disco, mostra a velocidade ao longo "
        "do dia, aponta excessos e paradas longas e guarda um laudo em PDF. "
        "Não substitui o disco físico nem um laudo oficial de oficina: ajuda a "
        "equipe a trabalhar mais rápido e com o mesmo critério em todo disco.",
        after=12,
    )

    # 2. É / não é / faz / não faz
    _titulo(doc, "2. É / não é  ·  Faz / não faz")
    _p(
        doc,
        "Esse quadro da Lean Inception evita expectativa errada. O lado esquerdo "
        "é o produto. O lado direito é o que fica de fora nesta versão do PGI.",
        after=10,
    )
    _subtitulo(doc, "O produto É  ×  NÃO É")
    _tabela_2x2(
        doc,
        ["É", "Não é"],
        [
            [
                "Um protótipo acadêmico utilizável (PGI)",
                "Um leitor de disco por visão computacional",
                "Um painel local (computador + navegador)",
                "Uma ajuda para auditoria interna da frota",
                "Um registro por placa (histórico + PDF)",
            ],
            [
                "Um app de celular (Android/iOS)",
                "Um sistema na nuvem / SaaS",
                "Um substituto jurídico do disco de papel",
                "Um laudo de oficina credenciada pelo Inmetro",
                "Um ERP de frota ou sistema de RH",
            ],
        ],
        ["D5F5E3", "FADBD8"],
    )
    _subtitulo(doc, "O produto FAZ  ×  NÃO FAZ")
    _tabela_2x2(
        doc,
        ["Faz", "Não faz"],
        [
            [
                "Recebe a foto do disco (PNG/JPG) e a placa",
                "Desenha a curva de velocidade das 24 horas",
                "Marca excesso acima de 80 km/h",
                "Marca parada longa (mais de 10 minutos parado)",
                "Estima a distância e guarda a viagem no banco",
                "Gera laudo PDF para arquivo interno",
                "Mostra erro claro se a foto estiver ilegível",
            ],
            [
                "Não faz login nem cadastro de usuário",
                "Não lê o nome manuscrito no miolo do disco (OCR)",
                "Não reconhece a placa na foto sozinho",
                "Não envia WhatsApp, e-mail automático nem pagamento",
                "Não conecta com sistema da transportadora",
                "Não fiscaliza o motorista em tempo real",
                "Não garante leitura perfeita em foto ruim ou torta",
            ],
        ],
        ["D5F5E3", "FADBD8"],
    )

    # 3. Objetivos
    _titulo(doc, "3. Os 3 objetivos principais")
    _p(
        doc,
        "Na Lean Inception, poucos objetivos bem claros valem mais do que uma "
        "lista enorme. Para o CronotaScan, neste PGI, são estes três:",
        after=10,
    )

    objetivos = [
        (
            "Objetivo 1 — Ler o disco de forma padronizada",
            "Tirar a conferência do “olhômetro”. A partir da foto, o sistema "
            "mostra a jornada em uma curva simples (hora × km/h), para qualquer "
            "pessoa da equipe ver o mesmo gráfico.",
        ),
        (
            "Objetivo 2 — Destacar o que importa na jornada",
            "Apontar sozinho os trechos acima de 80 km/h e as paradas longas "
            "(mais de 10 minutos). Assim o gestor não precisa caçar o traço "
            "no papel circular.",
        ),
        (
            "Objetivo 3 — Deixar rastro para auditoria",
            "Gravar cada leitura por placa e emitir um PDF com resumo, imagem "
            "desdobrada e alertas. O disco físico continua existindo; o sistema "
            "só organiza a evidência.",
        ),
    ]
    for titulo, texto in objetivos:
        _subtitulo(doc, titulo)
        _p(doc, texto, after=8)

    _p(
        doc,
        "Como saber se deu certo (bem simples): (1) uma foto boa gera curva visível; "
        "(2) excessos e paradas aparecem na tela; (3) dá para reabrir a viagem no "
        "histórico e baixar o PDF.",
        after=16,
    )

    _p(doc, "Referência do método", size=10, bold=True, after=2)
    _p(
        doc,
        "CAROLI, Paulo. Lean Inception: como alinhar pessoas e construir o produto certo. "
        "São Paulo: Editora Caroli, 2018. Artefatos usados: visão do produto; é / não é; "
        "faz / não faz; objetivos do produto.",
        size=10,
        after=4,
    )

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent / "Lean_Inception_CronotaScan.docx"
    montar(destino)
    print(destino)
