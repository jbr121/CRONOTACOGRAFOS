"""Gera a Definição do Projeto Integrador — CronotaScan (PGI)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1B, 0x4F, 0x72)
PRETO = RGBColor(0x1C, 0x28, 0x33)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _font(run, size=12, bold=False, italic=False, color=PRETO) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _sombra(cell, hex_cor: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_cor)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _bordas(table) -> None:
    tblPr = table._tbl.tblPr if table._tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "1B4F72")
        borders.append(el)
    tblPr.append(borders)


def _p(doc, texto, *, size=12, bold=False, italic=False, center=False, after=8, before=0) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 1.15
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    _font(run, size=size, bold=bold, italic=italic)


def _titulo(doc, texto) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(texto)
    _font(run, size=15, bold=True, color=AZUL)


def _item(doc, texto) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.clear()
    run = p.add_run(texto)
    _font(run, size=12)


def _tabela(doc, cabecalho: list[str], linhas: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordas(table)
    for i, titulo in enumerate(cabecalho):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(titulo)
        _font(run, size=11, bold=True, color=BRANCO)
        _sombra(cell, "1B4F72")
    for r, linha in enumerate(linhas):
        for c, valor in enumerate(linha):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(valor)
            _font(run, size=11)
            if r % 2 == 1:
                _sombra(cell, "EAF2F8")
    doc.add_paragraph()


def montar(saida: Path) -> None:
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    _p(doc, "UNIVERSIDADE DO VALE DO ITAJAÍ — UNIVALI", size=12, bold=True, center=True, after=2)
    _p(doc, "Curso de Sistemas para Internet", size=12, center=True, after=2)
    _p(doc, "Projeto Integrador (PGI)", size=12, center=True, after=14)
    _p(doc, "Definição do Projeto Integrador", size=20, bold=True, center=True, after=4)
    _p(doc, "CronotaScan — leitor de discos de cronotacógrafo analógico", size=13, italic=True, center=True, after=8)
    _p(doc, "José Eduardo  ·  Márcio Fossa", size=12, center=True, after=2)
    _p(doc, "Itajaí — SC", size=12, center=True, after=16)

    _titulo(doc, "1. Identificação")
    _tabela(
        doc,
        ["Item", "Descrição"],
        [
            ["Nome do produto", "CronotaScan"],
            ["Tipo", "Protótipo de software web local (PGI)"],
            ["Curso", "Sistemas para Internet — Univali"],
            ["Disciplina", "Projeto Integrador"],
            ["Integrantes", "José Eduardo e Márcio Fossa"],
            ["Repositório", "https://github.com/jbr121/CRONOTACOGRAFOS"],
        ],
    )

    _titulo(doc, "2. O que é o projeto")
    _p(
        doc,
        "O CronotaScan é um sistema que lê a foto de um disco de cronotacógrafo "
        "(o papel circular de 24 horas usado em caminhões e ônibus) e mostra, "
        "na tela, a velocidade ao longo do dia, os excessos, as paradas longas "
        "e um laudo em PDF. A pessoa informa a placa, envia a imagem e recebe "
        "o resultado no computador, sem precisar de internet na nuvem.",
    )

    _titulo(doc, "3. Problema")
    _p(
        doc,
        "Hoje a conferência do disco ainda é feita no olho: o operador olha o "
        "papel, estima a altura do traço e anota excessos e paradas à mão. "
        "Isso demora, cada pessoa lê de um jeito e fica difícil guardar prova "
        "padronizada por veículo. Quem sofre com isso é o operador (cansaço e "
        "retrabalho), o gestor de frota (falta de histórico claro) e o auditor "
        "interno (laudo irregular).",
    )

    _titulo(doc, "4. Justificativa")
    _p(
        doc,
        "O cronotacógrafo é equipamento obrigatório em vários veículos de carga "
        "e de passageiros. Mesmo com o modelo digital, muita frota ainda usa o "
        "disco de papel. Um protótipo acadêmico que padroniza a leitura ajuda "
        "a treinar a equipe, reduz erro de interpretação e gera um arquivo "
        "simples (histórico + PDF) para a auditoria interna. O projeto cabe no "
        "PGI porque une programação web, banco de dados e visão computacional "
        "em um fluxo único, possível de demonstrar na banca.",
    )

    _titulo(doc, "5. Objetivo geral")
    _p(
        doc,
        "Desenvolver um protótipo que auxilie a leitura de discos analógicos de "
        "cronotacógrafo, transformando a foto do disco em curva de velocidade, "
        "alertas e laudo, de forma padronizada e fácil de usar.",
    )

    _titulo(doc, "6. Objetivos específicos")
    _item(doc, "Permitir o envio da foto do disco e da placa do veículo.")
    _item(doc, "Gerar a curva de velocidade das 24 horas a partir da imagem.")
    _item(doc, "Detectar excesso acima de 80 km/h e parada longa (mais de 10 minutos).")
    _item(doc, "Guardar as leituras no computador (histórico por placa).")
    _item(doc, "Emitir um PDF com resumo, imagem desdobrada e alertas.")

    _titulo(doc, "7. Público-alvo")
    _item(doc, "Operador de leitura: tira a foto ou o scan e quer um resultado rápido.")
    _item(doc, "Gestor de frota: consulta o histórico por placa e vê excessos e paradas.")
    _item(doc, "Auditor interno: precisa de um PDF para arquivar junto do disco físico.")
    _p(
        doc,
        "Uso previsto: um computador Windows no setor de frota, na oficina ou "
        "no laboratório da Univali. Não é um aplicativo de celular.",
    )

    _titulo(doc, "8. Escopo (o que entra nesta versão)")
    _item(doc, "Upload de PNG/JPG + placa obrigatória.")
    _item(doc, "Leitura da coroa de velocidade e extração da curva.")
    _item(doc, "Alertas de excesso e de parada prolongada.")
    _item(doc, "Estimativa de distância.")
    _item(doc, "Dashboard, histórico e laudo PDF.")
    _item(doc, "Mensagem de erro clara se a foto não der para ler.")

    _titulo(doc, "9. Fora do escopo (o que não entra)")
    _item(doc, "Login, cadastro de usuário e sistema na nuvem.")
    _item(doc, "Aplicativo móvel.")
    _item(doc, "Leitura automática do nome escrito à mão no centro do disco.")
    _item(doc, "Integração com WhatsApp, ERP da transportadora ou pagamento.")
    _item(doc, "Substituir o disco físico ou o laudo oficial de oficina credenciada.")

    _titulo(doc, "10. Resultados esperados")
    _item(doc, "Software funcionando no navegador (endereço local da máquina).")
    _item(doc, "Uma foto boa gerar curva, números e PDF.")
    _item(doc, "Histórico consultável por placa.")
    _item(doc, "Documentação simples de uso para a banca.")
    _p(
        doc,
        "O sistema é uma ajuda. Foto escura, torta ou amassada pode errar. "
        "A decisão final continua sendo da pessoa que confere o disco.",
    )

    _titulo(doc, "11. Como o projeto será feito (resumo)")
    _p(
        doc,
        "O trabalho segue o fluxo do PGI: entender o problema, definir o que "
        "entra e o que não entra, construir o protótipo por partes e testar com "
        "imagens de disco. A parte visual da leitura fica em um módulo; a tela, "
        "o banco e o PDF ficam em outros. A equipe se organiza no GitHub, com "
        "reuniões curtas e revisão cruzada do código.",
    )

    _p(doc, "Itajaí, 2026.", italic=True, after=4, before=12)

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent / "Definicao_Projeto_Integrador_CronotaScan.docx"
    montar(destino)
    print(destino)
