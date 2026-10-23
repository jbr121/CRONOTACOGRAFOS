"""
Gera o documento Word com a seção de uso de IA no projeto CronotaScan.
Pode ser entregue separado ou colado no relatório M2.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


SAIDA = Path(__file__).resolve().parent / "CronotaScan_Uso_de_IA.docx"


def paragrafo(doc: Document, texto: str, negrito: bool = False, tamanho: int = 12) -> None:
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negrito
    run.font.size = Pt(tamanho)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def titulo_secao(doc: Document, texto: str) -> None:
    h = doc.add_heading(texto, level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)


def item_lista(doc: Document, texto: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(texto).font.size = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def main() -> None:
    doc = Document()

    # Margens ABNT: 3 cm esquerda/superior, 2 cm direita/inferior
    for section in doc.sections:
        section.top_margin    = Pt(85)   # ~3 cm
        section.bottom_margin = Pt(57)   # ~2 cm
        section.left_margin   = Pt(85)
        section.right_margin  = Pt(57)

    # ── Capa / cabeçalho ────────────────────────────────────────────────────
    titulo = doc.add_heading("Declaração de Uso de Inteligência Artificial", 0)
    titulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(6)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Projeto Integrador — CronotaScan")
    run.font.size = Pt(13)
    run.bold = True

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub2.add_run("Disciplina: Engenharia de Software II  |  Universidade do Vale do Itajaí — UNIVALI").font.size = Pt(11)

    sub3 = doc.add_paragraph()
    sub3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub3.add_run("Grupo: José Eduardo, Márcio Fossa, Erick").font.size = Pt(11)

    doc.add_paragraph()

    # ── 1. Introdução ────────────────────────────────────────────────────────
    titulo_secao(doc, "1. Introdução")
    paragrafo(doc,
        "Em conformidade com a política de transparência acadêmica da UNIVALI, "
        "este documento descreve os pontos do projeto CronotaScan em que foram "
        "utilizadas ferramentas de Inteligência Artificial Generativa (IAG) como "
        "suporte ao desenvolvimento. O uso de IA não substituiu o raciocínio crítico "
        "da equipe: todas as sugestões geradas foram revisadas, adaptadas e validadas "
        "pelos integrantes do grupo antes de serem incorporadas ao projeto."
    )

    # ── 2. Ferramenta utilizada ──────────────────────────────────────────────
    titulo_secao(doc, "2. Ferramenta de IA Utilizada")
    paragrafo(doc,
        "A ferramenta utilizada foi o Cursor (https://cursor.com), um ambiente de "
        "desenvolvimento integrado com IA baseado no modelo Claude (Anthropic). "
        "Dentro do Cursor, utilizamos o modo Cloud Agent, que permite ao assistente "
        "ler e editar arquivos do projeto de forma autônoma sob supervisão do desenvolvedor."
    )

    # ── 3. Onde foi utilizada ────────────────────────────────────────────────
    titulo_secao(doc, "3. Onde a IA Foi Utilizada no Projeto")

    doc.add_heading("3.1 Pipeline de Visão Computacional (core/processador_imagem.py)", level=3)
    paragrafo(doc,
        "A parte mais complexa do sistema é o pipeline OpenCV que transforma a foto "
        "do disco analógico em uma curva de velocidade × tempo. A IA foi utilizada para:"
    )
    item_lista(doc, "Sugerir o uso de cv2.warpPolar para o desdobramento polar do disco.")
    item_lista(doc, "Propor o uso de cv2.fitEllipse para detectar a borda do disco mesmo em fotos inclinadas.")
    item_lista(doc, "Auxiliar na lógica de homografia (correção de perspectiva elipse → círculo).")
    item_lista(doc, "Sugerir o uso de máscara HSV para isolar o papel branco do disco e ignorar o mousepad colorido de fundo.")
    item_lista(doc, "Depurar erros de detecção falsa de velocidade causados pelo resíduo da grade impressa.")
    paragrafo(doc,
        "Em todos esses casos, a equipe revisou o código gerado, ajustou os parâmetros "
        "com base nas fotos reais dos discos (ex.: faixas HSV do verde impresso, raios "
        "fracionários da escala de velocidade) e validou os resultados manualmente."
    )

    doc.add_heading("3.2 Detecção e Correção de Bugs", level=3)
    paragrafo(doc,
        "Durante o desenvolvimento, vários bugs foram identificados com o auxílio da IA:"
    )
    item_lista(doc, "Bug de fase: a madrugada parada era exibida como período de movimento — corrigido pela IA ao propor o uso da curva de velocidade extraída como score de orientação, em vez da máscara de atividade crua.")
    item_lista(doc, "Bug do mousepad: o fundo azul/escuro da mesa era detectado como o disco — corrigido com máscara de brancura (brilho − saturação) em vez de Hough simples.")
    item_lista(doc, "Contagem excessiva de quilômetros: a IA identificou que o zigue-zague do odômetro estava contando ruído e sugeriu filtro de distância mínima entre cruzamentos.")

    doc.add_heading("3.3 Testes Unitários (tests/test_pipeline.py)", level=3)
    paragrafo(doc,
        "Os 14 testes sintéticos do arquivo test_pipeline.py foram elaborados com "
        "suporte da IA, que gerou discos sintéticos com OpenCV para simular cenários "
        "como: fundo preto, mousepad azul, disco inclinado, madrugada parada e saltos "
        "irrealistas de velocidade. A equipe definiu os cenários e os critérios de aceite; "
        "a IA traduziu esses cenários em código de teste."
    )

    doc.add_heading("3.4 Documentação Acadêmica", level=3)
    paragrafo(doc,
        "A IA auxiliou na redação e formatação dos seguintes documentos entregues "
        "na disciplina de Engenharia de Software II:"
    )
    item_lista(doc, "M1 — Visão, Arquitetura e Planejamento do CronotaScan (Word + diagramas UML).")
    item_lista(doc, "Lean Inception — Visão do produto, É/Não é/Faz/Não faz e objetivos principais.")
    item_lista(doc, "Definição do Projeto Integrador (PGI) — escopo, cronograma e riscos.")
    item_lista(doc, "Atividade de Testes (partes 1 e 2) — respostas às questões teóricas sobre testes de software.")
    item_lista(doc, "M2 — Relatório de progresso, roteiro de apresentação e print das telas.")
    paragrafo(doc,
        "Em todos os casos, a equipe forneceu os dados reais do projeto (arquivos Python, "
        "telas da aplicação, nomes dos integrantes) e a IA organizou o conteúdo no "
        "formato solicitado pela professora. O texto foi revisado e aprovado pelo grupo "
        "antes da entrega."
    )

    doc.add_heading("3.5 Análise e Depuração de Desempenho", level=3)
    paragrafo(doc,
        "A IA foi consultada para analisar gargalos de desempenho, como a função "
        "_endireitar_grade_no_linear que levava 20 segundos por imagem usando "
        "cross-correlação pixel a pixel. A IA propôs substituir o algoritmo O(W×H×shifts) "
        "por centroide por coluna O(W×H), reduzindo o tempo para menos de 1 segundo."
    )

    # ── 4. O que NÃO foi feito pela IA ──────────────────────────────────────
    titulo_secao(doc, "4. O Que NÃO Foi Feito pela IA")
    paragrafo(doc,
        "Para clareza e integridade acadêmica, listamos abaixo as atividades que "
        "foram realizadas integralmente pela equipe, sem suporte de IA:"
    )
    item_lista(doc, "Coleta e digitalização das fotos reais de discos de cronotacógrafo.")
    item_lista(doc, "Definição dos requisitos funcionais do sistema junto ao orientador/professor.")
    item_lista(doc, "Validação dos resultados do pipeline contra laudos reais (comparação manual dos valores de km e velocidade).")
    item_lista(doc, "Apresentação oral do projeto para a banca.")
    item_lista(doc, "Decisões de arquitetura (escolha de FastAPI, SQLite, SPA HTML/JS puro).")

    # ── 5. Conclusão ─────────────────────────────────────────────────────────
    titulo_secao(doc, "5. Considerações Finais")
    paragrafo(doc,
        "O uso de IA generativa no projeto CronotaScan foi transparente, supervisionado "
        "e complementar ao trabalho da equipe. A ferramenta funcionou como um assistente "
        "de programação avançado: acelerou a escrita de código repetitivo, sugeriu "
        "abordagens algorítmicas e ajudou a formatar documentos, mas todas as decisões "
        "técnicas e acadêmicas relevantes foram tomadas pelos integrantes do grupo. "
        "A equipe se responsabiliza integralmente pelo conteúdo entregue."
    )

    doc.save(str(SAIDA))
    print(f"Documento salvo em: {SAIDA}")


if __name__ == "__main__":
    main()
