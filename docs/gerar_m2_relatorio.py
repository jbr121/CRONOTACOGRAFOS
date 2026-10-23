"""Gera o relatório M2 do CronotaScan (andamento das sprints, GC e mudanças)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1B, 0x4F, 0x72)
PRETO = RGBColor(0x1C, 0x28, 0x33)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _font(run, size: int = 12, bold: bool = False, italic: bool = False, color=PRETO) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _p(doc: Document, texto: str, *, size: int = 12, bold: bool = False,
       italic: bool = False, align: str = "justify", after: int = 8, before: int = 0,
       first: bool = True) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.first_line_indent = Cm(1.25) if first and align == "justify" else Cm(0)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    run = p.add_run(texto)
    _font(run, size=size, bold=bold, italic=italic)


def _titulo(doc: Document, texto: str, nivel: int = 1) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(18 if nivel == 1 else 12)
    pf.space_after = Pt(8)
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto)
    _font(run, size=14 if nivel == 1 else 12, bold=True, color=AZUL)


def _sombra(cell, cor_hex: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), cor_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _bordas(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "1B4F72")
        borders.append(el)
    tblPr.append(borders)


def _tabela(doc: Document, cabecalho: list[str], linhas: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordas(table)
    # Cabeçalho
    hdr_cells = table.rows[0].cells
    for i, txt in enumerate(cabecalho):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(txt)
        _font(run, size=11, bold=True, color=BRANCO)
        _sombra(hdr_cells[i], "1B4F72")
    # Linhas
    for r, linha in enumerate(linhas):
        for c, valor in enumerate(linha):
            cell = table.rows[r + 1].cells[c]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(valor)
            _font(run, size=11)
            if r % 2 == 1:
                _sombra(cell, "EAF2F8")
    doc.add_paragraph()


def _capa(doc: Document) -> None:
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(3.0)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)

    _p(doc, "UNIVERSIDADE DO VALE DO ITAJAÍ — UNIVALI", size=12, bold=True, align="center", first=False, after=2)
    _p(doc, "Curso de Sistemas para Internet", size=12, align="center", first=False, after=2)
    _p(doc, "Disciplina: Engenharia de Software II", size=12, align="center", first=False, after=8)
    _p(doc, "M2 — Andamento das Sprints, Gerência de Configuração e de Mudanças", size=16, bold=True,
       align="center", first=False, after=6)
    _p(doc, "Projeto: CronotaScan — leitor de discos de cronotacógrafo analógico", size=13, italic=True,
       align="center", first=False, after=18)
    _p(doc, "Integrantes: José Eduardo · Márcio Fossa · Erick", size=12, align="center", first=False, after=2)
    _p(doc, "Repositório: https://github.com/jbr121/CRONOTACOGRAFOS", size=11, align="center", first=False, after=18)
    _p(doc, "Itajaí — SC", size=12, align="center", first=False, after=2)
    _p(doc, "Outubro de 2026", size=12, align="center", first=False, after=2)
    doc.add_page_break()


def montar_relatorio(saida: Path) -> None:
    doc = Document()
    _capa(doc)

    _titulo(doc, "1. Objetivo", nivel=1)
    _p(
        doc,
        "Este relatório apresenta o andamento do projeto CronotaScan desde o início das sprints da "
        "disciplina de Engenharia de Software II. O foco é mostrar o que foi planejado, o que foi "
        "entregue, o que ainda falta, como as mudanças de escopo foram tratadas e como estamos "
        "controlando o repositório e a organização da equipe. Sempre que possível, cada afirmação "
        "é acompanhada de evidências (prints, links de commits, pull requests e quadro de tarefas).",
    )

    _titulo(doc, "3.1 Andamento das sprints", nivel=1)
    _p(
        doc,
        "Até esta etapa realizamos pelo menos duas sprints. Para cada uma, o quadro a seguir deve ser "
        "preenchido com os itens planejados, o que foi concluído e os principais problemas. "
        "As células marcadas com colchetes devem ser completadas pelo grupo com as informações reais.",
    )
    _tabela(
        doc,
        ["Sprint / período", "Objetivo", "Planejado (backlog)", "Concluído", "Não concluído / motivo", "Incremento"],
        [
            [
                "Sprint 1 — [datas]",
                "Colocar o fluxo básico no ar (upload da imagem, curva de velocidade e histórico inicial).",
                "Ex.: configurar FastAPI; tela de upload; leitura do disco sintético; configurar GitHub Projects.",
                "Ex.: upload funcionando; curva no navegador; banco SQLite criado.",
                "Ex.: foto real pendente — problema com centro/mousepad.",
                "Protótipo básico processando um disco sintético.",
            ],
            [
                "Sprint 2 — [datas]",
                "Melhorar a leitura da foto real e gerar o laudo PDF.",
                "Ex.: centro pela grade verde; correção de perspectiva; laudo PDF; ajustes no dashboard.",
                "Itens que o grupo finalizou de fato.",
                "Itens que ficaram para a próxima sprint e o porquê.",
                "Processar foto real de referência e emitir laudo PDF.",
            ],
        ],
    )
    _p(
        doc,
        "Além dos quadros acima, o relatório deve anexar prints do quadro de tarefas no início e no fim de "
        "cada sprint, bem como, se houver, o gráfico de burndown gerado pela ferramenta de gestão.",
    )

    _titulo(doc, "3.2 Resultado das sprints – incremento funcional", nivel=1)
    _p(
        doc,
        "Até o momento, o CronotaScan evoluiu de testes sintéticos para fotos mais próximas da realidade. "
        "Hoje o sistema permite enviar placa e imagem, detectar o centro do disco usando papel e grade verde, "
        "corrigir a perspectiva da foto, desdobrar a coroa de velocidade, detectar excessos e paradas, gravar "
        "a viagem no histórico e gerar um laudo em PDF. Nesta seção, o grupo deve relacionar os incrementos "
        "entregues com os itens do backlog e com os requisitos atendidos.",
    )
    _tabela(
        doc,
        ["Requisito / US", "Sprint", "Evidência (print / link)"],
        [
            ["RF01/RF02 – upload + placa", "Sprint 1", "[link do commit ou PR; print da tela de upload]"],
            ["RF03/RF04 – centro e perspectiva", "Sprint 2", "[evidência da foto real processada]"],
            ["RF14 – laudo PDF", "Sprint 2", "[print do PDF aberto e link do PR]"],
        ],
    )
    _p(
        doc,
        "Também é importante listar o que ainda falta para as próximas sprints (funcionalidades pendentes e "
        "ajustes de qualidade), deixando claro que o planejamento continua vivo.",
    )

    _titulo(doc, "3.3 Gerenciamento de mudanças", nivel=1)
    _p(
        doc,
        "As principais mudanças de escopo, requisitos ou arquitetura desde a M1 foram registradas no quadro "
        "abaixo. Cada linha corresponde a uma mudança aprovada, com análise de impacto e justificativa.",
    )
    _tabela(
        doc,
        ["ID", "Solicitação", "Tipo", "Impacto / análise", "Decisão", "Alterações e justificativa"],
        [
            [
                "M01",
                "Ajustar RF03 para usar papel + grade verde como referência de centro.",
                "Alteração de requisito",
                "Centro só pelo furo falhava em fotos com mousepad; o polar escapava para fora do disco.",
                "Aprovada na Sprint 1.",
                "Atualizamos a especificação e o código de detecção, adicionando testes sintéticos. "
                "Backlog recebeu uma US específica de robustez de centro.",
            ],
            [
                "M02",
                "Incluir laudo PDF como entrega obrigatória.",
                "Entrada de escopo",
                "Aumenta esforço mas aumenta muito o valor para auditoria interna.",
                "Aprovada na Sprint 2.",
                "Criamos o componente GeradorLaudoPDF, ajustamos requisitos e backlog e incluímos a rota de PDF.",
            ],
            [
                "M03",
                "Refinar forma de calcular a distância (uso do odômetro gráfico).",
                "Alteração de requisito",
                "Apenas integrar a curva inflava a quilometragem em algumas fotos reais.",
                "Aprovada em reunião de revisão.",
                "Alteramos RF09 para priorizar o zigue‑zague do odômetro; o analisador passou a expor origem_distancia.",
            ],
        ],
    )

    _titulo(doc, "3.4 Rastreabilidade", nivel=1)
    _p(
        doc,
        "A matriz de rastreabilidade foi atualizada para refletir as mudanças registradas na seção anterior. "
        "Cada requisito funcional principal está ligado ao caso de uso correspondente, ao modelo (por exemplo, "
        "diagrama de classes), à classe de implementação e ao commit ou PR onde aquele requisito foi tratado.",
    )

    _titulo(doc, "3.5 Gerenciamento de configuração", nivel=1)
    _p(
        doc,
        "O repositório CronotaScan está organizado em camadas: core/ concentra a lógica de visão computacional, "
        "regras de negócio e geração de PDF; main.py expõe as rotas FastAPI; static/ contém a SPA; docs/ guarda "
        "os documentos do PGI/ES2; tests/ reúne os testes automatizados. A estratégia de branches segue o "
        "GitHub Flow: a branch main guarda versões estáveis e cada incremento é desenvolvido em uma branch de "
        "feature (por exemplo, feat/odometro), integrada via pull request.",
    )
    _p(
        doc,
        "Adotamos mensagens de commit curtas e descritivas, muitas vezes com prefixos como feat:, fix: e docs:. "
        "Sempre que possível, um integrante abre o PR e outro faz a revisão, registrando comentários. A partir "
        "da versão 0.9.x passamos a marcar releases com tags (v0.9.8, v0.9.9, etc.), facilitando a referência "
        "nas apresentações e nos laudos.",
    )

    _titulo(doc, "3.6 Comunicação e organização da equipe", nivel=1)
    _p(
        doc,
        "A comunicação diária do grupo é feita sobretudo pelo WhatsApp, onde combinamos prazos, horários de "
        "reunião e divisão de tarefas. As decisões técnicas mais importantes (por exemplo, aceitar uma mudança "
        "de requisito ou o resultado de um teste com foto real) são registradas em comentários de pull request "
        "no GitHub, para ficarem ligadas diretamente ao código.",
    )
    _p(
        doc,
        "Utilizamos um quadro de tarefas (GitHub Projects ou Trello) com colunas A Fazer, Fazendo, Em Teste e "
        "Pronto. Cada item tem um responsável (José, Márcio ou Erick), o que ajuda na cobrança e na evidência "
        "de contribuição individual.",
    )

    _titulo(doc, "3.7 Contribuição individual", nivel=1)
    _tabela(
        doc,
        ["Integrante", "Principais tarefas na M2", "Commits / PRs (exemplos)"],
        [
            [
                "José Eduardo",
                "Ajustes no processador de imagem (centro, perspectiva, odômetro), testes sintéticos e "
                "roteiros de verificação.",
                "[links de commits e PRs do José]",
            ],
            [
                "Márcio Fossa",
                "Rotas FastAPI, ajustes na SPA, geração de laudo PDF e documentação de uso.",
                "[links de commits e PRs do Márcio]",
            ],
            [
                "Erick",
                "Tarefas de apoio definidas pelo grupo (por exemplo, testes manuais, organização do "
                "quadro, documentação ou código específico).",
                "[links de commits e PRs do Erick]",
            ],
        ],
    )
    _p(
        doc,
        "Esta tabela deve ser complementada com os links reais de commits, PRs e outras evidências produzidas "
        "por cada integrante durante a M2, conforme o Anexo D do enunciado.",
    )

    _titulo(doc, "3.8 Conhecimento técnico e justificativa das escolhas", nivel=1)
    _p(
        doc,
        "Linguagem e framework: escolhemos Python com FastAPI porque a disciplina já utiliza o ecossistema "
        "Python, a integração com OpenCV é direta e a aplicação é leve para rodar nos laboratórios da Univali. "
        "Para o banco de dados optamos por SQLite, pois atende bem ao escopo local do PGI (um arquivo por "
        "projeto, sem necessidade de servidor dedicado).",
    )
    _p(
        doc,
        "A arquitetura em camadas (SPA, API, domínio, dados) foi adotada para separar bem responsabilidades: "
        "a SPA trata de interação com o usuário; a API, das rotas e erros HTTP; o domínio, da visão "
        "computacional e das regras de negócio; e a camada de dados, da persistência e evidências. Essa "
        "separação facilita manutenção, testes unitários e a explicação do projeto na banca.",
    )
    _p(
        doc,
        "As principais bibliotecas escolhidas foram OpenCV (visão computacional), SQLAlchemy (ORM sobre SQLite) "
        "e ReportLab (laudos PDF). Suas vantagens são ampla documentação, uso consolidado em projetos similares "
        "e boa integração com o restante do código Python. Como limitação, o uso intensivo de OpenCV exige mais "
        "cuidado com performance em fotos de alta resolução, e o SQLite não é indicado para uso concorrente em "
        "produção, mas atende perfeitamente ao cenário local do PGI.",
    )

    _p(
        doc,
        "Em resumo, as decisões técnicas foram guiadas por simplicidade, facilidade de uso em laboratório e "
        "coerência com os requisitos do projeto, sempre com atenção à rastreabilidade entre backlog, código "
        "e evidências apresentadas neste relatório.",
        after=4,
    )

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


def montar_roteiro(saida: Path) -> None:
    """Roteiro de fala para a apresentação oral (bullet points)."""
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    _p(doc, "Roteiro da apresentação M2 — CronotaScan", size=16, bold=True, align="center", first=False, after=6)
    _p(doc, "Resumo em tópicos para o grupo ler antes de apresentar.", size=11, italic=True,
       align="center", first=False, after=12)

    _titulo(doc, "Slide 1 – Visão geral")
    _p(doc, "• Apresentar o grupo (José, Márcio e Erick) e o projeto CronotaScan.", first=False)
    _p(doc, "• Explicar em uma frase: 'leitor de discos de cronotacógrafo a partir da foto'.", first=False)
    _p(doc, "• Dizer que a M2 mostra como as sprints andaram, como controlamos mudanças e o repositório.", first=False)

    _titulo(doc, "Slide 2 – Objetivo das sprints")
    _p(doc, "• Sprint 1: deixar o fluxo básico pronto (upload, curva, histórico).", first=False)
    _p(doc, "• Sprint 2: foto real mais robusta e laudo PDF.", first=False)

    _titulo(doc, "Slide 3 – Andamento da Sprint 1")
    _p(doc, "• Mostrar print do quadro no início e no fim.", first=False)
    _p(doc, "• Falar o que estava planejado e o que foi concluído.", first=False)
    _p(doc, "• Dizer qual foi o incremento: 'já conseguimos processar um disco sintético'.", first=False)

    _titulo(doc, "Slide 4 – Andamento da Sprint 2")
    _p(doc, "• Mesmo esquema: planejado × realizado.", first=False)
    _p(doc, "• Incremento: 'agora processa foto real e gera PDF'.", first=False)

    _titulo(doc, "Slide 5 – Demonstração")
    _p(doc, "• Combinar antes quem vai operar: enviar foto, mostrar dashboard, histórico e PDF.", first=False)

    _titulo(doc, "Slide 6 – Mudanças de escopo")
    _p(doc, "• Destacar 2 ou 3 mudanças importantes do quadro de mudanças (M01–M03).", first=False)
    _p(doc, "• Explicar rapidamente por que cada mudança foi feita.", first=False)

    _titulo(doc, "Slide 7 – Gerência de configuração")
    _p(doc, "• Mostrar a estrutura de pastas do GitHub.", first=False)
    _p(doc, "• Explicar uso de branches de feature + PR + revisão.", first=False)
    _p(doc, "• Apontar histórico de commits e, se tiver, tags de versão.", first=False)

    _titulo(doc, "Slide 8 – Comunicação e quadro de tarefas")
    _p(doc, "• Dizer que o grupo usa WhatsApp para o dia a dia.", first=False)
    _p(doc, "• Mostrar o quadro de tarefas com responsáveis.", first=False)

    _titulo(doc, "Slide 9 – Contribuição de cada integrante")
    _p(doc, "• José fala o que fez (visão computacional, testes, etc.).", first=False)
    _p(doc, "• Márcio fala o que fez (API, UI, PDF, docs).", first=False)
    _p(doc, "• Erick fala o que fez (tarefas que vocês definirem).", first=False)

    _titulo(doc, "Slide 10 – Escolhas técnicas e próximos passos")
    _p(doc, "• Falar de forma simples: Python + FastAPI + OpenCV + SQLite.", first=False)
    _p(doc, "• Dizer por que isso funciona bem para o PGI.", first=False)
    _p(doc, "• Encerrar com 'o que falta' até a entrega final.", first=False)

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


if __name__ == "__main__":
    pasta = Path(__file__).resolve().parent
    montar_relatorio(pasta / "M2_Relatorio_CronotaScan.docx")
    montar_roteiro(pasta / "M2_Roteiro_Apresentacao_CronotaScan.docx")
    print(pasta / "M2_Relatorio_CronotaScan.docx")
    print(pasta / "M2_Roteiro_Apresentacao_CronotaScan.docx")

