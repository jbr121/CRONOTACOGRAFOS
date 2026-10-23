"""Respostas da atividade de testes — linguagem simples."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1B, 0x4F, 0x72)
PRETO = RGBColor(0x1C, 0x28, 0x33)


def _font(run, size=12, bold=False, italic=False, color=PRETO) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _p(doc, texto, *, size=12, bold=False, italic=False, center=False, after=8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    _font(run, size=size, bold=bold, italic=italic)


def _titulo(doc, texto) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texto)
    _font(run, size=13, bold=True, color=AZUL)


def montar(saida: Path) -> None:
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    _p(doc, "Atividade — Testes de software", size=18, bold=True, center=True, after=4)
    _p(doc, "Respostas em linguagem simples", size=12, italic=True, center=True, after=14)

    _titulo(doc, "1. Dá para achar todos os defeitos de um software? Justifique.")
    _p(
        doc,
        "Não. Ninguém consegue garantir que um programa está 100% sem erro. "
        "O teste serve para achar problemas, não para provar que não existe nenhum. "
        "Se fôssemos testar todas as combinações possíveis (todos os cliques, "
        "todos os números, todos os caminhos), o trabalho nunca acabaria. "
        "Por isso o teste diminui o risco, mas não deixa o software perfeito.",
    )

    _titulo(doc, "2. Quais são as condições mínimas para um bom teste?")
    _p(
        doc,
        "Para testar de verdade, e não só “mexer no sistema”, precisa no mínimo: "
        "saber o que o programa deveria fazer; ter casos de teste com entrada e "
        "resultado esperado; um ambiente parecido com o real; anotar o que passou "
        "e o que falhou; e combinar quando o teste termina. Sem isso, a equipe "
        "fica só clicando sem critério.",
    )

    _titulo(doc, "3. Cite dois critérios para encerrar os testes.")
    _p(
        doc,
        "Dois critérios bem usados no plano de testes: primeiro, ter rodado o "
        "que estava planejado (por exemplo, todos os casos combinados, ou os "
        "requisitos mais importantes cobertos). Segundo, a qualidade estar "
        "aceitável: nenhum erro grave aberto. Acabar só porque o prazo acabou "
        "também acontece, mas é o pior motivo para parar.",
    )

    _titulo(doc, "4. Qual o ponto fraco de testar só se a função do sistema funciona?")
    _p(
        doc,
        "Se a gente testa só o “caminho feliz” da tela, pode deixar passar muita "
        "coisa: o sistema lento, falha de segurança, tela confusa, módulo interno "
        "quebrado ou duas partes que não conversam bem. Ou seja: a função até "
        "aparece certa para o usuário, mas o software ainda pode falhar na vida real.",
    )

    _titulo(doc, "5. Quantos testes seriam necessários para achar a maioria dos defeitos?")
    _p(
        doc,
        "Não existe um número pronto, tipo “50 testes e acabou”. Um teste bom é "
        "aquele que acha um erro que ninguém tinha visto. Na prática, poucos "
        "testes bem pensados (valores no limite, fluxos que o usuário mais usa, "
        "casos que costumam quebrar) encontram a maior parte dos problemas. "
        "Os erros que sobram pedem muito mais tempo. Por isso a gente prioriza "
        "o que tem mais risco, em vez de tentar testar tudo.",
    )

    _titulo(doc, "6. Qual percentual do esforço deve ir para os testes?")
    _p(
        doc,
        "Nos livros de engenharia de software, em projetos feitos por etapas "
        "(modelo em cascata ou evolutivo), o teste costuma ficar com cerca de "
        "30% a 40% do trabalho. Se o sistema for crítico, esse número sobe. "
        "O recado é simples: teste não é “se sobrar tempo no final”.",
    )

    _titulo(doc, "7. Qual o problema de o programador testar o que ele mesmo fez?")
    _p(
        doc,
        "Quem escreveu o código quer ver ele funcionando. Acaba testando do "
        "mesmo jeito que programou e não enxerga o próprio erro. Falta um olhar "
        "de fora. Por isso o desenvolvedor pode (e deve) fazer um primeiro teste, "
        "mas outra pessoa precisa olhar também, senão muita falha passa.",
    )

    _titulo(doc, "8. Qual a diferença entre teste de integração e teste de sistema?")
    _p(
        doc,
        "O teste de integração pergunta: as peças, depois de prontas, se encaixam? "
        "Ele olha a ligação entre os módulos (se um chama o outro certo, se os "
        "dados passam direito). O teste de sistema pergunta: o software inteiro "
        "atende o que foi pedido? Aí se testa o produto completo, como o usuário "
        "vai usar, e também coisas como lentidão e segurança.",
    )

    _titulo(doc, "9. O que é integração top-down e bottom-up?")
    _p(
        doc,
        "Top-down é de cima para baixo. Começa pelas telas e pelo controle geral. "
        "O que ainda não existe embaixo é trocado por um “faz de conta” (stub). "
        "A vantagem é ver o sistema cedo. A desvantagem é que as funções de "
        "baixo, muitas vezes as mais importantes, só são testadas no fim.",
    )
    _p(
        doc,
        "Bottom-up é de baixo para cima. Começa pelos cálculos e pelo banco, "
        "e sobe. Usa um programinha que chama o módulo (driver). A base fica "
        "mais segura, mas a tela e o sistema como um todo só aparecem no final. "
        "Muita equipe mistura as duas formas.",
    )

    _titulo(doc, "10. Por que o teste de sistema deve ser feito por uma equipe independente?")
    _p(
        doc,
        "Porque quem construiu o sistema está tendencioso: conhece o atalho, "
        "sabe “como era para usar” e defende o próprio trabalho. Uma equipe de "
        "fora testa como o usuário testa e tenta achar furo, não proteger o código. "
        "Isso não tira o teste que o programador faz no pedaço dele; só completa, "
        "principalmente quando o sistema já está montado.",
    )

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent / "Respostas_Atividade_Testes_simples.docx"
    montar(destino)
    print(destino)
